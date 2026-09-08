from __future__ import annotations

from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook

from investment_analyzer.connectors.models import Source
from investment_analyzer.connectors.seed import ensure_default_sources
from investment_analyzer.db import create_all_tables, create_db_engine, create_session_factory
from investment_analyzer.entity_resolution.models import Entity, IdentifierType
from investment_analyzer.entity_resolution.service import IdentifierSpec, find_or_create_entity
from investment_analyzer.fundamentals.metrics import Metric
from investment_analyzer.normalization.models import DataPoint, ValueKind
from investment_analyzer.reports.bundle import build_report_bundle
from investment_analyzer.reports.excel_export import (
    _sanitize_excel_string,
    build_excel_workbook,
    export_excel,
)
from investment_analyzer.reports.json_export import report_bundle_to_dict

FETCHED_AT = datetime(2024, 3, 1, tzinfo=UTC)

EXPECTED_SHEETS = [
    "Zusammenfassung", "Kennzahlen", "Bewertung", "Risiken", "Nachrichten", "Quellen", "Annahmen",
]


def _session_factory(tmp_path: Path):
    engine = create_db_engine(f"sqlite:///{tmp_path / 'reports-excel-test.db'}")
    create_all_tables(engine)
    return create_session_factory(engine)


def _dp(
    *, entity: Entity, source: Source, metric: Metric, period_end: date, value: float,
    retrieved_at_utc: datetime = FETCHED_AT,
) -> DataPoint:
    return DataPoint(
        entity_id=entity.id,
        source_id=source.id,
        metric_name=metric.value,
        period_start=None,
        period_end=period_end,
        published_at=period_end,
        retrieved_at_utc=retrieved_at_utc,
        value_raw=str(value),
        value_normalized=value,
        unit="USD",
        currency="USD",
        value_kind=ValueKind.REPORTED,
        document_url="https://example.invalid/doc",
        document_type="10-K",
        content_hash="x" * 64,
        document_id=f"acc-{metric.value}-{period_end.isoformat()}",
    )


def _build_bundle(session, tmp_path: Path):
    sources = ensure_default_sources(session)
    entity = find_or_create_entity(
        session, name="Firma G (synthetisches Beispiel)",
        identifiers=[IdentifierSpec(id_type=IdentifierType.CIK, id_value="0000000007")],
    )
    session.flush()

    for d, u in zip(
        [date(2020, 12, 31), date(2021, 12, 31), date(2022, 12, 31), date(2023, 12, 31)],
        [1000.0, 1100.0, 1210.0, 1331.0],
        strict=True,
    ):
        session.add(_dp(entity=entity, source=sources["sec_edgar"], metric=Metric.REVENUE, period_end=d, value=u))
        session.add(
            _dp(entity=entity, source=sources["sec_edgar"], metric=Metric.OPERATING_CASH_FLOW, period_end=d, value=u * 0.15)
        )
        session.add(_dp(entity=entity, source=sources["sec_edgar"], metric=Metric.CAPEX, period_end=d, value=-u * 0.05))

    session.add(
        _dp(
            entity=entity, source=sources["sec_edgar"], metric=Metric.PRICE_CLOSE, period_end=date(2024, 1, 15),
            value=20.0, retrieved_at_utc=datetime(2024, 1, 15, tzinfo=UTC),
        )
    )
    session.commit()
    return build_report_bundle(session, entity, as_of=datetime(2024, 3, 1, tzinfo=UTC))


def test_build_excel_workbook_hat_alle_sieben_tabellenblaetter(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        wb = build_excel_workbook(bundle)

    assert wb.sheetnames == EXPECTED_SHEETS


def test_export_excel_kann_geladen_werden(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        out_path = tmp_path / "report.xlsx"
        export_excel(bundle, out_path)

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert wb.sheetnames == EXPECTED_SHEETS


def test_zusammenfassung_zeigt_dieselben_werte_wie_json_export(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        wb = build_excel_workbook(bundle)
        as_dict = report_bundle_to_dict(bundle)

    ws = wb["Zusammenfassung"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 9)}
    assert values["Unternehmen"] == as_dict["header"]["entity_name"]
    assert values["Gesamtscore (0-100)"] == as_dict["score"]["total_score"]


def test_zusammenfassung_enthaelt_pflichthinweis(tmp_path: Path) -> None:
    """Auftrag §12: Hinweis muss an jeder Berichtsausgabe sichtbar sein."""

    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        wb = build_excel_workbook(bundle)

    ws = wb["Zusammenfassung"]
    all_rows = [
        (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value) for r in range(1, ws.max_row + 1)
    ]
    hinweis_werte = [value for label, value in all_rows if label == "Hinweis"]
    assert hinweis_werte
    assert "keine Anlageberatung" in hinweis_werte[0]


def test_quellen_sheet_enthaelt_alle_registrierten_quellen(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        wb = build_excel_workbook(bundle)

    ws = wb["Quellen"]
    keys = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert keys == {"sec_edgar", "alpha_vantage", "gdelt", "ir_rss"}


def test_annahmen_sheet_enthaelt_roic_hinweis(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        bundle = _build_bundle(session, tmp_path)
        wb = build_excel_workbook(bundle)

    ws = wb["Annahmen"]
    all_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any("ROIC-Steuersatz" in (v or "") for v in all_values)


def test_sanitize_excel_string_neutralisiert_formel_trigger() -> None:
    assert _sanitize_excel_string("=SUM(A1:A2)") == "'=SUM(A1:A2)"
    assert _sanitize_excel_string("+1234") == "'+1234"
    assert _sanitize_excel_string("-1234") == "'-1234"
    assert _sanitize_excel_string("@SUM") == "'@SUM"
    assert _sanitize_excel_string("Normaler Text") == "Normaler Text"


def test_leerer_bundle_erzeugt_dennoch_gueltige_arbeitsmappe(tmp_path: Path) -> None:
    session_factory = _session_factory(tmp_path)
    with session_factory() as session:
        ensure_default_sources(session)
        entity = find_or_create_entity(
            session, name="Firma ohne Daten",
            identifiers=[IdentifierSpec(id_type=IdentifierType.CIK, id_value="0000000099")],
        )
        session.commit()
        bundle = build_report_bundle(session, entity)
        wb = build_excel_workbook(bundle)

    assert wb.sheetnames == EXPECTED_SHEETS
